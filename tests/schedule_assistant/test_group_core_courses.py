import datetime as dtm
import io

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import GradientFill, PatternFill

from src.schedule_assistant.core_courses.location_parser import Item
from src.schedule_assistant.core_courses.parser import subject_cell_fill_color
from src.schedule_assistant.modules.parser.core_courses_adapter import (
    _expand_nested_location_lessons,
    expand_grouped_core_course_lessons,
    group_core_course_lessons,
    grouped_core_course_to_json,
)
from src.schedule_assistant.modules.parser.schemas import Lesson


def _lesson(
    *,
    lesson_name: str = "History",
    course_name: str = "BS - Year 1",
    weekday: str = "FRIDAY",
    start_time: dtm.time = dtm.time(9, 0),
    end_time: dtm.time = dtm.time(10, 30),
    lesson_class_type: str | None = "lec",
    modifiers: Item | None = None,
    date_on: list[dtm.date] | None = None,
    date_except: list[dtm.date] | None = None,
    date_from: dtm.date | None = None,
    color: str | None = None,
) -> Lesson:
    return Lesson(
        lesson_name=lesson_name,
        color=color,
        lesson_class_type=lesson_class_type,
        weekday=weekday,
        start_time=start_time,
        end_time=end_time,
        course_name=course_name,
        group_name=("B25-CSE-01",),
        teacher="Teacher A",
        room="107",
        source_type="core_course",
        date_on=date_on,
        date_except=date_except,
        date_from=date_from,
        spreadsheet_id="sheet-id",
        google_sheet_gid="gid",
        google_sheet_name="SUMMER",
        modifiers=modifiers,
    )


TERM_DATES = {
    "SUMMER": (
        dtm.date(2026, 6, 1),
        dtm.date(2026, 8, 2),
    ),
}


def test_group_merges_same_subject_slots_into_one_entry():
    lessons = [
        _lesson(start_time=dtm.time(9, 0)),
        _lesson(start_time=dtm.time(10, 40), end_time=dtm.time(12, 10)),
    ]
    grouped = group_core_course_lessons(lessons, term_dates=TERM_DATES)

    assert len(grouped) == 1
    assert grouped[0].cohort == "BS - Year 1"
    assert grouped[0].subject == "History"
    assert grouped[0].start_date == dtm.date(2026, 6, 1)
    assert grouped[0].end_date == dtm.date(2026, 8, 2)
    assert len(grouped[0].components) == 2
    assert grouped[0].components[0].type == "lec"
    assert grouped[0].components[1].start_time == dtm.time(10, 40)


def test_subject_cell_fill_color_handles_merged_and_missing_fills() -> None:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet["B2"].fill = PatternFill("solid", fgColor="FFa1b2c3")
    sheet.merge_cells("B2:C2")
    sheet["D2"].fill = PatternFill("solid", fgColor="80112233")
    sheet["E2"].fill = GradientFill(stop=("112233", "445566"))
    sheet["G2"].fill = PatternFill("solid", fgColor="00abcdef")
    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    reloaded = load_workbook(buffer).active
    assert reloaded is not None

    assert subject_cell_fill_color(reloaded, "B2") == "#A1B2C3"
    assert subject_cell_fill_color(reloaded, "C2") == "#A1B2C3"
    assert subject_cell_fill_color(reloaded, "D2") is None
    assert subject_cell_fill_color(reloaded, "E2") is None
    assert subject_cell_fill_color(reloaded, "F2") is None
    assert subject_cell_fill_color(reloaded, "G2") == "#ABCDEF"


def test_group_uses_majority_non_null_course_color() -> None:
    grouped = group_core_course_lessons(
        [
            _lesson(color="#112233"),
            _lesson(start_time=dtm.time(10, 40), end_time=dtm.time(12, 10), color="#112233"),
            _lesson(start_time=dtm.time(12, 40), end_time=dtm.time(14, 10), color="#AABBCC"),
            _lesson(start_time=dtm.time(14, 20), end_time=dtm.time(15, 50)),
        ],
        term_dates=TERM_DATES,
    )

    assert grouped[0].color == "#112233"


def test_group_rejects_tied_course_colors() -> None:
    with pytest.raises(ValueError, match=r"History.*#112233.*#AABBCC"):
        group_core_course_lessons(
            [
                _lesson(color="#112233"),
                _lesson(start_time=dtm.time(10, 40), end_time=dtm.time(12, 10), color="#AABBCC"),
            ],
            term_dates=TERM_DATES,
        )


def test_group_splits_different_subjects_and_cohorts():
    lessons = [
        _lesson(lesson_name="History", course_name="BS - Year 1"),
        _lesson(lesson_name="Math", course_name="BS - Year 2"),
    ]
    grouped = group_core_course_lessons(lessons, term_dates=TERM_DATES)

    assert len(grouped) == 2
    subjects = {entry.subject for entry in grouped}
    assert subjects == {"History", "Math"}


def test_group_omits_location_only_modifiers():
    lessons = [_lesson(modifiers=Item(location="108"))]
    grouped = group_core_course_lessons(lessons, term_dates=TERM_DATES)

    assert grouped[0].components[0].modifiers is None


def test_group_preserves_modifiers_on_components():
    modifiers = Item(location="460", except_=[dtm.date(2026, 7, 17)])
    lessons = [
        _lesson(
            modifiers=modifiers,
            date_except=[dtm.date(2026, 7, 17)],
        )
    ]
    grouped = group_core_course_lessons(lessons, term_dates=TERM_DATES)
    component = grouped[0].components[0]

    assert component.modifiers is not None
    assert component.modifiers.location == "460"
    assert component.modifiers.except_ == [dtm.date(2026, 7, 17)]
    dumped = grouped_core_course_to_json(grouped[0])["components"][0]
    assert "date_on" not in dumped
    assert "date_except" not in dumped
    assert "date_from" not in dumped
    assert dumped["modifiers"]["except_"] == ["2026-07-17"]


def test_group_single_component_per_slot_when_modifiers_have_nest():
    nested = Item(location="106", on=[dtm.date(2026, 7, 17)])
    root = Item(location="107", NEST=[nested])
    lessons = [
        _lesson(modifiers=root, start_time=dtm.time(9, 0)),
        _lesson(
            modifiers=root.model_copy(deep=True),
            start_time=dtm.time(10, 40),
            end_time=dtm.time(12, 10),
        ),
    ]
    grouped = group_core_course_lessons(lessons, term_dates=TERM_DATES)

    assert len(grouped) == 1
    assert len(grouped[0].components) == 2
    for component in grouped[0].components:
        assert component.room == "107"
        assert component.modifiers is not None
        assert component.modifiers.NEST is not None
        assert component.modifiers.NEST[0].location == "106"


def test_expand_nested_location_lessons_for_collisions():
    nested = Item(location="106", on=[dtm.date(2026, 7, 17)])
    root = Item(location="107", NEST=[nested])
    lessons = [_lesson(modifiers=root)]
    expanded = _expand_nested_location_lessons(lessons)

    assert len(expanded) == 2
    assert expanded[0].room == "107"
    assert expanded[0].date_except == [dtm.date(2026, 7, 17)]
    assert expanded[1].room == "106"
    assert expanded[1].date_on == [dtm.date(2026, 7, 17)]


def test_expand_group_round_trip():
    modifiers = Item(location="ONLINE", on=[dtm.date(2026, 6, 13)])
    lessons = [
        _lesson(
            lesson_name="История России",
            course_name="BS - Year 1",
            modifiers=modifiers,
            date_on=[dtm.date(2026, 6, 13)],
        ),
    ]
    grouped = group_core_course_lessons(lessons, term_dates=TERM_DATES)
    expanded = expand_grouped_core_course_lessons(grouped)

    assert len(expanded) == 1
    assert expanded[0].lesson_name == "История России"
    assert expanded[0].course_name == "BS - Year 1"
    assert expanded[0].lesson_class_type == "lec"
    assert expanded[0].room == "107"
    assert expanded[0].modifiers is not None
    assert expanded[0].modifiers.location == "ONLINE"
    assert expanded[0].date_on == [dtm.date(2026, 6, 13)]
