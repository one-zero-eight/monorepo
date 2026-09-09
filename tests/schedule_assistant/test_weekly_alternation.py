import datetime as dtm
import io
from collections import Counter
from typing import Any

import icalendar
import pytest
from httpx import AsyncClient
from openpyxl import load_workbook
from pydantic import ValidationError

from src.schedule_assistant.modules.bookings.client import BookingDTO
from src.schedule_assistant.modules.bookings.match import (
    booking_coverage,
    find_extra_auto_bookings,
    find_matching_auto_booking,
    iter_booking_occurrences,
    iter_payload_occurrences,
    payload_matches_auto_booking,
)
from src.schedule_assistant.modules.bookings.review import (
    build_review_index,
    collect_booking_payloads,
    split_payloads_around_conflicts,
)
from src.schedule_assistant.modules.issues.booking_match import slot_has_matching_booking
from src.schedule_assistant.modules.issues.booking_slots import build_bookable_slots
from src.schedule_assistant.modules.issues.checker import IssueChecker
from src.schedule_assistant.modules.issues.per_week import per_week_issues_from_schedule_config
from src.schedule_assistant.modules.issues.placement import iter_concrete_dates, meetings_overlap
from src.schedule_assistant.modules.issues.schemas import OccurrencePlacement, ScheduledMeeting, WeeklyPatternPlacement
from src.schedule_assistant.modules.schedule.domain import meetings_from_schedule_config
from src.schedule_assistant.modules.schedule.export_xlsx import expand_meetings, export_schedule_xlsx
from src.schedule_assistant.modules.schedule.ics import render_calendar
from src.schedule_assistant.modules.schedule.service import _meeting_identity, _teacher_weekly_meetings
from src.schedule_assistant.modules.schedule_config.instructor_meetings import count_meetings_by_instructor
from src.schedule_assistant.modules.schedule_config.repository import ScheduleConfigRepository
from src.schedule_assistant.modules.schedule_config.schemas import (
    ComponentSessionSeries,
    CourseConfig,
    CoursesConfig,
    ScheduleConfig,
    SectionsConfig,
    TermConfig,
    WeeklyAlternation,
    WeeklyPatternSlot,
)
from src.schedule_assistant.modules.schedule_config.weekly_dates import (
    active_weekly_dates,
    expand_weekly_slot,
    normalize_alternation,
)
from src.schedule_assistant.modules.schedule_config.weekly_pattern_canonicalization import canonicalize_weekly_slot
from src.schedule_assistant.weekday import Weekday


def config(
    *, weekday: str = "TUESDAY", anchor: str | None = "2026-09-16", edits: list[dict[str, Any]] | None = None
) -> ScheduleConfig:
    return ScheduleConfig.model_validate(
        {
            "term": {
                "name": "Fall",
                "semester": {"start_date": "2026-09-02", "end_date": "2026-10-31"},
                "sections": [
                    {"code": "core", "name": "Core", "programs": [{"code": "BS", "name": "BS", "groups": ["G1"]}]}
                ],
            },
            "rooms": [{"id": "108", "name": "108"}],
            "instructors": [{"id": "teacher@innopolis.ru"}, {"id": "other@innopolis.ru"}],
            "students_groups": [{"code": "G1"}],
            "courses": [
                {
                    "name": "Algebra",
                    "section_code": "core",
                    "components": [
                        {
                            "tag": "lec",
                            "per_week": 1,
                            "audience": ["G1"],
                            "sessions": [
                                {
                                    "weekly_pattern": [
                                        {
                                            "weekday": weekday,
                                            "start_time": "09:00",
                                            "end_time": "10:30",
                                            "room": "108",
                                            "instructor": "teacher@innopolis.ru",
                                            "alternation": {"anchor_week": anchor} if anchor else None,
                                            "edits": edits,
                                        }
                                    ],
                                }
                            ],
                        }
                    ],
                }
            ],
        }
    )


def slot_for(value: ScheduleConfig) -> WeeklyPatternSlot:
    sessions = value.courses[0].components[0].sessions
    assert sessions and sessions[0].weekly_pattern
    return sessions[0].weekly_pattern[0]


def meetings(value: ScheduleConfig, *, expand: bool = True) -> list[ScheduledMeeting]:
    return meetings_from_schedule_config(
        CoursesConfig(courses=value.courses),
        SectionsConfig(sections=value.term.sections),
        value.term,
        expand_weekly=expand,
    )


def date_strings(value: ScheduleConfig) -> list[str]:
    return [
        meeting.placement.date.isoformat()
        for meeting in meetings(value)
        if isinstance(meeting.placement, OccurrencePlacement)
    ]


def test_contract_requires_anchor_and_preserves_null() -> None:
    assert (
        WeeklyPatternSlot.model_validate({"weekday": "MONDAY", "start_time": "09:00", "end_time": "10:30"}).alternation
        is None
    )
    with pytest.raises(ValidationError):
        WeeklyAlternation.model_validate({})
    with pytest.raises(ValidationError):
        WeeklyPatternPlacement.model_validate({"weekday": "MONDAY", "alternation": {}})
    assert WeeklyPatternPlacement.model_validate(
        {"weekday": "MONDAY", "alternation": {"anchor_week": "2026-09-16"}}
    ).alternation == WeeklyAlternation(anchor_week=dtm.date(2026, 9, 16))


@pytest.mark.parametrize(
    ("anchor", "expected"),
    [
        ("2026-09-16", ["2026-09-15", "2026-09-29", "2026-10-13", "2026-10-27"]),
        ("2026-09-09", ["2026-09-08", "2026-09-22", "2026-10-06", "2026-10-20"]),
        (
            None,
            [
                "2026-09-08",
                "2026-09-15",
                "2026-09-22",
                "2026-09-29",
                "2026-10-06",
                "2026-10-13",
                "2026-10-20",
                "2026-10-27",
            ],
        ),
    ],
)
def test_partial_first_week_and_both_phases(anchor, expected) -> None:
    assert date_strings(config(anchor=anchor)) == expected


def test_year_boundary_non_monday_start_and_anchor_both_directions() -> None:
    window = TermConfig.DateRange(start_date=dtm.date(2026, 12, 27), end_date=dtm.date(2027, 1, 25))
    alternation = WeeklyAlternation(anchor_week=dtm.date(2027, 1, 11))
    assert normalize_alternation(alternation, Weekday.WEDNESDAY) == WeeklyAlternation(anchor_week=dtm.date(2027, 1, 6))
    assert active_weekly_dates(window, Weekday.MONDAY, Weekday.WEDNESDAY, alternation) == [
        dtm.date(2026, 12, 28),
        dtm.date(2027, 1, 11),
        dtm.date(2027, 1, 25),
    ]
    assert active_weekly_dates(window, Weekday.TUESDAY, Weekday.WEDNESDAY, alternation) == [
        dtm.date(2026, 12, 29),
        dtm.date(2027, 1, 12),
    ]
    assert active_weekly_dates(window, Weekday.WEDNESDAY, Weekday.WEDNESDAY, alternation) == [
        dtm.date(2027, 1, 6),
        dtm.date(2027, 1, 20),
    ]


def test_audience_window_clips_without_resetting_phase() -> None:
    value = config()
    value.term.sections[0].programs[0].semester = TermConfig.DateRange(
        start_date=dtm.date(2026, 9, 20), end_date=dtm.date(2026, 10, 20)
    )
    assert date_strings(value) == ["2026-09-29", "2026-10-13"]


def test_inactive_edits_stay_dormant_and_active_moves_apply_all_fields() -> None:
    value = config(
        edits=[
            {"select_week": "2026-09-08", "date": "2026-09-16", "instructor": "other@innopolis.ru"},
            {"select_week": "2026-09-15", "cancel": True},
            {
                "select_week": "2026-09-29",
                "date": "2026-10-06",
                "start_time": "12:40",
                "end_time": "14:10",
                "room": "109",
                "instructor": "other@innopolis.ru",
            },
        ]
    )
    assert date_strings(value) == ["2026-10-06", "2026-10-13", "2026-10-27"]
    moved = meetings(value)[0]
    assert (moved.start_time, moved.end_time, moved.room, moved.instructor) == (
        dtm.time(12, 40),
        dtm.time(14, 10),
        "109",
        "other@innopolis.ru",
    )
    assert count_meetings_by_instructor(value.courses, value.term, ["teacher@innopolis.ru", "other@innopolis.ru"]) == {
        "teacher@innopolis.ru": 2,
        "other@innopolis.ru": 1,
    }
    placement = meetings(value, expand=False)[0]
    assert list(iter_concrete_dates(placement, start_date=dtm.date(2026, 10, 6), end_date=dtm.date(2026, 10, 6))) == [
        dtm.date(2026, 10, 6)
    ]


def test_opposite_phases_and_disjoint_windows_do_not_conflict_but_moves_do() -> None:
    first = meetings(config(), expand=False)[0]
    opposite = meetings(config(anchor="2026-09-09"), expand=False)[0]
    assert not meetings_overlap(first, opposite)
    assert meetings_overlap(first, meetings(config(), expand=False)[0])
    assert meetings_overlap(first, meetings(config(anchor=None), expand=False)[0])
    other_window = opposite.model_copy(deep=True)
    assert isinstance(other_window.placement, WeeklyPatternPlacement)
    other_window.placement.start_date = dtm.date(2027, 1, 1)
    other_window.placement.end_date = dtm.date(2027, 2, 1)
    assert not meetings_overlap(first, other_window)
    moved = meetings(
        config(edits=[{"select_week": "2026-09-15", "date": "2026-09-08", "start_time": "12:40", "end_time": "14:10"}]),
        expand=False,
    )[0]
    assert not meetings_overlap(moved, opposite)
    opposite.start_time, opposite.end_time = dtm.time(12, 40), dtm.time(14, 10)
    assert meetings_overlap(moved, opposite)


def test_canonicalization_preserves_phase_inactive_edits_and_actual_events() -> None:
    value = config(
        edits=[
            {"select_week": date, "room": "109"} for date in ["2026-09-15", "2026-09-29", "2026-10-13", "2026-10-27"]
        ]
        + [{"select_week": "2026-09-08", "date": "2026-09-10", "room": "110"}]
    )
    slot = slot_for(value)
    canonical = canonicalize_weekly_slot(slot, value.term, ["G1"])
    assert canonical.alternation == WeeklyAlternation(anchor_week=dtm.date(2026, 9, 14))
    assert canonical.room == "109"
    assert slot.edits
    assert canonical.edits == [slot.edits[-1]]
    before = [item.occurrence for item in expand_weekly_slot(slot, value.term.semester, value.term.starting_day)]
    after = [item.occurrence for item in expand_weekly_slot(canonical, value.term.semester, value.term.starting_day)]
    assert before == after


def decoded_ics_dates(data: bytes) -> Counter[tuple[dtm.date, dtm.time, dtm.time]]:
    result: Counter[tuple[dtm.date, dtm.time, dtm.time]] = Counter()
    events = icalendar.Calendar.from_ical(data).walk("VEVENT")
    overrides = {
        (str(event["uid"]), event.decoded("recurrence-id")): event for event in events if "recurrence-id" in event
    }
    for event in events:
        if "recurrence-id" in event:
            continue
        start, end = event.decoded("dtstart"), event.decoded("dtend")
        if "rrule" not in event:
            result[(start.date(), start.time(), end.time())] += 1
            continue
        rule = event["rrule"]
        assert isinstance(rule, icalendar.vRecur)
        step = dtm.timedelta(weeks=int(rule.get("INTERVAL", [1])[0]))
        until = rule["UNTIL"][0]
        exclusions = event.get("exdate", [])
        if not isinstance(exclusions, list):
            exclusions = [exclusions]
        excluded = {item.dt for exclusion in exclusions for item in exclusion.dts}
        while start <= until:
            override = overrides.get((str(event["uid"]), start))
            if override:
                actual_start, actual_end = override.decoded("dtstart"), override.decoded("dtend")
                result[(actual_start.date(), actual_start.time(), actual_end.time())] += 1
            elif start not in excluded:
                result[(start.date(), start.time(), end.time())] += 1
            start, end = start + step, end + step
    return result


@pytest.mark.parametrize("anchor", ["2026-09-09", "2026-09-16", None])
def test_domain_booking_ics_and_xlsx_have_identical_actual_dates(anchor) -> None:
    value = config(
        anchor=anchor,
        edits=[
            {"select_week": "2026-09-08", "cancel": True},
            {"select_week": "2026-09-15", "date": "2026-09-16", "start_time": "12:40", "end_time": "14:10"},
            {"select_week": "2026-09-29", "cancel": True},
        ],
    )
    expected = Counter(
        (meeting.placement.date, meeting.start_time, meeting.end_time)
        for meeting in meetings(value)
        if isinstance(meeting.placement, OccurrencePlacement)
    )
    slots = build_bookable_slots(
        CoursesConfig(courses=value.courses), SectionsConfig(sections=value.term.sections), value.term, {"108"}
    )
    actual_booking = Counter(
        (start.date(), start.time(), end.time())
        for slot in slots
        for start, end in iter_payload_occurrences(slot.payload)
    )
    assert actual_booking == expected
    assert Counter((meeting.date, meeting.start, meeting.end) for meeting in expand_meetings(value)) == expected
    data = render_calendar("Algebra", [("G1", meeting) for meeting in meetings(value, expand=False)])
    assert decoded_ics_dates(data) == expected
    for slot in slots:
        if slot.payload.get("recurrence"):
            assert slot.payload["recurrence"].get("interval", 1) == (2 if anchor else 1)
    if anchor:
        for layout in ("groups", "compact_groups"):
            value.term.sections[0].default_layout = layout
            book = load_workbook(io.BytesIO(export_schedule_xlsx(value)[0]), rich_text=True)
            assert any("Every other week" in str(cell.value) for sheet in book for row in sheet for cell in row)


def test_ics_uid_distinguishes_phases() -> None:
    uids = []
    for anchor in ("2026-09-09", "2026-09-16"):
        data = render_calendar("Algebra", [("G1", meetings(config(anchor=anchor), expand=False)[0])])
        event = icalendar.Calendar.from_ical(data).walk("VEVENT")[0]
        rule = event["rrule"]
        assert isinstance(rule, icalendar.vRecur)
        assert rule["INTERVAL"] == [2]
        uids.append(str(event["uid"]))
    assert uids[0] != uids[1]


def test_booking_interval_phase_extra_detection_and_conflict_splitting() -> None:
    value = config()
    slot = build_bookable_slots(
        CoursesConfig(courses=value.courses), SectionsConfig(sections=value.term.sections), value.term, {"108"}
    )[0]
    original = iter_payload_occurrences(slot.payload)
    split = split_payloads_around_conflicts(slot, [(original[1][0], original[1][1], [])])
    assert {occurrence for part in split for occurrence in iter_payload_occurrences(part)} == set(original) - {
        original[1]
    }
    assert all(part["recurrence"]["interval"] == 2 for part in split)
    own: dict[str, Any] = {
        **slot.payload,
        "title": "Auto: Algebra (lec)",
        "uid": "own",
        "outlook_booking_id": "master",
        "recurrence_complete": True,
    }
    assert payload_matches_auto_booking(slot.payload, own)
    assert slot_has_matching_booking(slot.payload, auto_bookings=[own], existing_bookings=[])
    wrong = {**own, "recurrence": {**own["recurrence"], "interval": 1}}
    assert not payload_matches_auto_booking(slot.payload, wrong)
    assert find_matching_auto_booking(slot.payload, [wrong]) is None
    assert not slot_has_matching_booking(slot.payload, auto_bookings=[wrong], existing_bookings=[])
    assert booking_coverage(slot.payload, [wrong]).uncertain
    assert find_extra_auto_bookings([wrong], [slot.payload]) == [wrong]
    opposite = {**own, "recurrence": {**own["recurrence"], "start_date": "2026-09-22"}}
    assert not payload_matches_auto_booking(slot.payload, opposite)
    assert not set(iter_booking_occurrences(opposite)) & set(original)
    recurrence = own["recurrence"]
    xml = f"<t:WeeklyRecurrence><t:Interval>2</t:Interval><t:DaysOfWeek>Tuesday</t:DaysOfWeek></t:WeeklyRecurrence><t:StartDate>{recurrence['start_date']}</t:StartDate><t:EndDate>{recurrence['until_date']}</t:EndDate>"
    assert iter_booking_occurrences({**own, "recurrence": xml}) == original
    weekly_xml = xml.replace("<t:Interval>2", "<t:Interval>1")
    dto = BookingDTO.model_validate(
        {**own, "recurrence": weekly_xml, "room_response": "Accept", "room_presence": "present", "can_cancel": True}
    )
    index = build_review_index([slot], auto_bookings=[dto], existing_bookings=[])
    assert not collect_booking_payloads(index, [slot.slot_id], {})
    assert index.extras and all(not extra["can_cancel"] for extra in index.extras.values())


@pytest.mark.asyncio
async def test_api_persists_normalized_anchor_phase_changes_and_disabling(
    authenticated_client: AsyncClient,
    schedule_config_repo: ScheduleConfigRepository,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr("src.schedule_assistant.dependencies.settings.moderator_emails", ["test@test.com"])
    value = config()
    response = await authenticated_client.put("/schedule-config/", json=value.model_dump(mode="json"))
    assert response.status_code == 200, response.text
    for anchor, normalized in [("2026-09-16", "2026-09-14"), ("2026-09-09", "2026-09-07"), (None, None)]:
        course = config(anchor=anchor).courses[0]
        response = await authenticated_client.put(
            "/schedule-config/courses/Algebra", json=course.model_dump(mode="json")
        )
        assert response.status_code == 200, response.text
        stored = schedule_config_repo.get_course("Algebra")
        assert stored and stored.components[0].sessions and stored.components[0].sessions[0].weekly_pattern
        alternation = stored.components[0].sessions[0].weekly_pattern[0].alternation
        assert (alternation.anchor_week.isoformat() if alternation else None) == normalized


@pytest.mark.parametrize("anchor", ["2026-09-09", "2026-09-16", None])
def test_room_teacher_group_student_conflicts_follow_actual_dates(anchor) -> None:
    first, second = config(), config(anchor=anchor)
    second.courses[0].name = "Other course"
    first.courses.extend(second.courses)
    first.students_groups[0].students = ["student@innopolis.ru"]
    sections = SectionsConfig(sections=first.term.sections, students_groups=first.students_groups)
    checker = IssueChecker(courses=CoursesConfig(courses=first.courses), sections=sections, term=first.term)
    concrete = meetings(first)
    has_conflict = anchor != "2026-09-09"
    assert bool(checker.check_for_room_issue(concrete)) == has_conflict
    assert bool(checker.check_for_teacher_issue(concrete)) == has_conflict
    assert bool(checker.check_for_group_issue(concrete)) == has_conflict
    sections.students_groups.append(first.students_groups[0].model_copy(update={"code": "G2"}))
    student_meetings = [
        item.model_copy(update={"groups": ("G2",)}) if item.course_name == "Other course" else item for item in concrete
    ]
    assert bool(checker.check_for_student_issue(student_meetings)) == has_conflict
    assert not per_week_issues_from_schedule_config(CoursesConfig(courses=first.courses), sections)


def test_calendar_xlsx_non_monday_weeks_keep_actual_dates() -> None:
    value = config(weekday="MONDAY", anchor="2027-01-11")
    value.term.semester = TermConfig.DateRange(start_date=dtm.date(2026, 12, 27), end_date=dtm.date(2027, 1, 25))
    value.term.starting_day = Weekday.WEDNESDAY
    value.term.sections[0].default_layout = "calendar"
    book = load_workbook(io.BytesIO(export_schedule_xlsx(value)[0]))
    dates = []
    for sheet in book:
        if sheet.cell(1, 2).value != "Monday":
            continue
        for row in sheet:
            if not str(row[0].value or "").startswith("Week "):
                continue
            next_row = row[0].row + 1
            if "Algebra" in str(sheet.cell(next_row, 2).value):
                dates.append(sheet.cell(row[0].row, 2).value)
    assert dates == ["December 28", "January 11", "January 25"]
    assert decoded_ics_dates(
        render_calendar("Algebra", [("G1", item) for item in meetings(value, expand=False)])
    ) == Counter((meeting.date, meeting.start, meeting.end) for meeting in expand_meetings(value))


def test_no_active_dates_does_not_emit_booking_or_ics_series() -> None:
    value = config()
    value.term.semester = TermConfig.DateRange(start_date=dtm.date(2026, 9, 2), end_date=dtm.date(2026, 9, 4))
    assert not date_strings(value)
    assert not build_bookable_slots(
        CoursesConfig(courses=value.courses), SectionsConfig(sections=value.term.sections), value.term, {"108"}
    )
    data = render_calendar("Algebra", [("G1", item) for item in meetings(value, expand=False)])
    assert not icalendar.Calendar.from_ical(data).walk("VEVENT")


def test_mfai_independent_lecture_tutorial_and_four_lab_phases() -> None:
    value = config()
    value.term.sections[0].programs[0].groups = [f"G{i}" for i in range(1, 5)]
    base = slot_for(value)
    value.courses[0].components = [
        CourseConfig.Component(tag="lec", audience=["@BS"], sessions=[ComponentSessionSeries(weekly_pattern=[base])]),
        CourseConfig.Component(
            tag="tut",
            audience=["@BS"],
            sessions=[ComponentSessionSeries(weekly_pattern=[base.model_copy(update={"weekday": Weekday.WEDNESDAY})])],
        ),
        CourseConfig.Component(
            tag="lab",
            per_group=True,
            audience=["@BS"],
            sessions=[
                ComponentSessionSeries(
                    audience=[f"G{i}"],
                    weekly_pattern=[
                        base.model_copy(
                            update={
                                "weekday": day,
                                "end_time": dtm.time(12, 10),
                                "room": str(108 + i),
                                "alternation": WeeklyAlternation(anchor_week=dtm.date(2026, 9, 7)),
                            }
                        )
                    ],
                )
                for i, day in enumerate([Weekday.TUESDAY, Weekday.WEDNESDAY, Weekday.THURSDAY, Weekday.FRIDAY], 1)
            ],
        ),
    ]
    concrete = meetings(value)
    teaching = [item for item in concrete if item.component_tag != "lab"]
    labs = [item for item in concrete if item.component_tag == "lab"]
    assert len(teaching) == 9 and len(labs) == 16
    assert {item.room for item in labs} == {"109", "110", "111", "112"}
    assert all(not meetings_overlap(first, second) for first in teaching for second in labs)
    assert all(item.end_time == dtm.time(12, 10) for item in labs)


def test_teacher_calendar_ignores_inactive_substitution_and_retains_phase_identity() -> None:
    value = config(
        edits=[
            {"select_week": "2026-09-08", "instructor": "other@innopolis.ru"},
            {"select_week": "2026-09-15", "instructor": "other@innopolis.ru", "date": "2026-09-16"},
        ]
    )
    meeting = meetings(value, expand=False)[0]
    substitutes = _teacher_weekly_meetings(meeting, {"other@innopolis.ru"})
    assert len(substitutes) == 1
    assert substitutes[0].placement == OccurrencePlacement(date=dtm.date(2026, 9, 16))
    regular = _teacher_weekly_meetings(meeting, {"teacher@innopolis.ru"})
    assert decoded_ics_dates(render_calendar("Regular", [("teacher", item) for item in regular])) == Counter(
        {
            (dtm.date(2026, 9, 29), dtm.time(9), dtm.time(10, 30)): 1,
            (dtm.date(2026, 10, 13), dtm.time(9), dtm.time(10, 30)): 1,
            (dtm.date(2026, 10, 27), dtm.time(9), dtm.time(10, 30)): 1,
        }
    )
    assert _meeting_identity(meeting) != _meeting_identity(meetings(config(anchor="2026-09-09"), expand=False)[0])
